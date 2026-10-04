"""Export surplus funds to CSV, Excel (.xlsx) or an HTML page."""

import csv
import html
from collections.abc import Callable, Iterable
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from surplus_funds.models import SurplusFund


def _local_time(value: datetime | None) -> datetime | None:
    """Timestamps are stored in UTC; show them in local time, without microseconds.

    Excel can't store time zones, so the result is a naive datetime.
    """
    if value is None:
        return None
    return value.astimezone().replace(tzinfo=None, microsecond=0)


# Header and value getter for every exported column, in order.
COLUMNS: list[tuple[str, Callable[[SurplusFund], Any]]] = [
    ("State", lambda f: f.source.state),
    ("County", lambda f: f.source.county),
    ("Parcel ID", lambda f: f.parcel_id),
    ("Case Number", lambda f: f.case_number),
    ("Owner", lambda f: f.owner_name),
    ("Buyer", lambda f: f.buyer_name),
    ("Property Address", lambda f: f.property_address),
    ("City", lambda f: f.city),
    ("Amount", lambda f: f.amount),
    ("Sale Date", lambda f: f.sale_date),
    ("Sale Date Precision", lambda f: f.sale_date_precision),
    ("Status", lambda f: f.status),
    ("First Seen", lambda f: _local_time(f.first_seen_at)),
    ("Last Seen", lambda f: _local_time(f.last_seen_at)),
    ("Source URL", lambda f: f.source.url),
]

HEADERS = [header for header, _ in COLUMNS]

# Excel number formats per column; other columns use the default.
_EXCEL_FORMATS = {
    "Amount": '"$"#,##0.00',
    "Sale Date": "yyyy-mm-dd",
    "First Seen": "yyyy-mm-dd hh:mm",
    "Last Seen": "yyyy-mm-dd hh:mm",
}
_MAX_COLUMN_WIDTH = 50


def _row(fund: SurplusFund) -> list[Any]:
    return [getter(fund) for _, getter in COLUMNS]


def export_csv(funds: Iterable[SurplusFund], path: Path) -> int:
    """Write funds to a CSV file. Returns the number of rows written."""
    count = 0
    # utf-8-sig adds a BOM, so Excel opens names with accents correctly.
    with open(path, "w", newline="", encoding="utf-8-sig") as file:
        writer = csv.writer(file)
        writer.writerow(HEADERS)
        for fund in funds:
            writer.writerow(_row(fund))
            count += 1
    return count


def export_excel(funds: Iterable[SurplusFund], path: Path) -> int:
    """Write funds to an .xlsx file. Returns the number of rows written."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Surplus Funds"

    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    count = 0
    for fund in funds:
        sheet.append(_row(fund))
        count += 1

    for index, header in enumerate(HEADERS, start=1):
        letter = get_column_letter(index)
        cells = sheet[letter]
        number_format = _EXCEL_FORMATS.get(header)
        if number_format:
            for cell in cells[1:]:
                cell.number_format = number_format
        width = max(len(str(cell.value)) for cell in cells if cell.value is not None)
        sheet.column_dimensions[letter].width = min(width + 2, _MAX_COLUMN_WIDTH)

    sheet.freeze_panes = "A2"  # header stays visible while scrolling
    sheet.auto_filter.ref = sheet.dimensions  # filter dropdowns on the header row
    workbook.save(path)
    return count


# The stylesheet is a separate file, not inline: Jenkins' default Content
# Security Policy blocks inline styles and scripts in published HTML reports.
_HTML_CSS = """\
body { font-family: Segoe UI, Helvetica, Arial, sans-serif; margin: 24px; color: #222; }
h1 { font-size: 22px; margin-bottom: 4px; }
.summary { color: #555; margin-bottom: 16px; }
table { border-collapse: collapse; font-size: 13px; }
th { background: #2f4f6f; color: #fff; text-align: left; position: sticky; top: 0; }
th, td { padding: 6px 10px; border-bottom: 1px solid #ddd; white-space: nowrap; }
tr:nth-child(even) td { background: #f5f7fa; }
td.number { text-align: right; font-variant-numeric: tabular-nums; }
"""


def _html_cell(header: str, value: Any) -> str:
    if value is None:
        return "<td></td>"
    if header == "Source URL":
        url = html.escape(value)
        return f'<td><a href="{url}">source</a></td>'
    if isinstance(value, Decimal):
        return f'<td class="number">${value:,.2f}</td>'
    if isinstance(value, datetime):
        return f"<td>{value:%Y-%m-%d %H:%M}</td>"
    if isinstance(value, date):
        return f"<td>{value:%Y-%m-%d}</td>"
    return f"<td>{html.escape(str(value))}</td>"


def export_html(funds: Iterable[SurplusFund], path: Path) -> int:
    """Write funds to an HTML page plus a stylesheet next to it. Returns the row count."""
    funds = list(funds)
    css_path = path.with_suffix(".css")
    css_path.write_text(_HTML_CSS, encoding="utf-8")

    total = sum((fund.amount for fund in funds), Decimal(0))
    generated = datetime.now().strftime("%Y-%m-%d %H:%M")
    header_row = "".join(f"<th>{html.escape(header)}</th>" for header in HEADERS)
    body_rows = "\n".join(
        "<tr>" + "".join(_html_cell(header, value) for header, value in zip(HEADERS, _row(fund), strict=True)) + "</tr>"
        for fund in funds
    )

    path.write_text(
        f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<title>Surplus Funds</title>
<link rel="stylesheet" href="{html.escape(css_path.name)}">
</head>
<body>
<h1>Surplus Funds</h1>
<p class="summary">{len(funds)} funds, total ${total:,.2f}. Generated {generated}.</p>
<table>
<thead><tr>{header_row}</tr></thead>
<tbody>
{body_rows}
</tbody>
</table>
</body>
</html>
""",
        encoding="utf-8",
    )
    return len(funds)


_EXPORTERS: dict[str, Callable[[Iterable[SurplusFund], Path], int]] = {
    ".csv": export_csv,
    ".xlsx": export_excel,
    ".html": export_html,
}


def check_export_path(path: Path) -> None:
    """Raise ValueError if the file extension isn't a supported format."""
    if path.suffix.lower() not in _EXPORTERS:
        raise ValueError(f"unsupported file type {path.suffix!r}, use {', '.join(_EXPORTERS)}")


def export_funds(funds: Iterable[SurplusFund], path: Path) -> int:
    """Export to CSV, Excel or HTML depending on the file extension."""
    check_export_path(path)
    return _EXPORTERS[path.suffix.lower()](funds, path)
