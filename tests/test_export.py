import csv
from datetime import date, datetime, timezone
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from surplus_funds.export import HEADERS, export_funds
from surplus_funds.models import Source, SurplusFund


@pytest.fixture
def funds() -> list[SurplusFund]:
    """Funds built in memory; export doesn't need the database."""
    source = Source(key="ga_hall", state="GA", county="Hall", url="https://example.com/list.pdf", file_format="pdf")
    seen = datetime(2026, 10, 4, 19, 24, 0, 3999, tzinfo=timezone.utc)
    return [
        SurplusFund(
            source=source,
            external_id="15025A000047|2016-11-25",
            parcel_id="15025A000047",
            owner_name="NEWCOMB ALVIN",
            buyer_name="JOSE MANZO-MADRIZAL",
            property_address="3442 CANDLER ROAD",
            city="GAINESVILLE",
            amount=Decimal("1220.17"),
            sale_date=date(2016, 11, 25),
            sale_date_precision="day",
            first_seen_at=seen,
            last_seen_at=seen,
        ),
        SurplusFund(
            source=source,
            external_id="X-1",
            parcel_id="X-1",
            owner_name="ŠIKIMIĆ ĐORĐE",
            amount=Decimal("5.00"),
            first_seen_at=seen,
            last_seen_at=seen,
        ),
    ]


def test_csv(funds, tmp_path):
    path = tmp_path / "funds.csv"
    assert export_funds(funds, path) == 2

    with open(path, newline="", encoding="utf-8-sig") as file:
        rows = list(csv.DictReader(file))

    assert list(rows[0]) == HEADERS
    assert rows[0]["Parcel ID"] == "15025A000047"
    assert rows[0]["Amount"] == "1220.17"
    assert rows[0]["Sale Date"] == "2016-11-25"
    assert rows[0]["Source URL"] == "https://example.com/list.pdf"
    assert "." not in rows[0]["First Seen"]  # no microseconds
    # Missing values are empty cells, accented names survive.
    assert rows[1]["Sale Date"] == ""
    assert rows[1]["Owner"] == "ŠIKIMIĆ ĐORĐE"


def test_excel(funds, tmp_path):
    path = tmp_path / "funds.xlsx"
    assert export_funds(funds, path) == 2

    sheet = load_workbook(path).active
    rows = list(sheet.iter_rows(values_only=True))
    assert list(rows[0]) == HEADERS
    assert len(rows) == 3

    record = dict(zip(HEADERS, rows[1], strict=True))
    assert record["Owner"] == "NEWCOMB ALVIN"
    # Real numbers and dates, not text, so Excel can sum and sort them.
    assert record["Amount"] == pytest.approx(1220.17)
    assert record["Sale Date"] == datetime(2016, 11, 25)
    assert isinstance(record["First Seen"], datetime)

    amount_cell = sheet.cell(row=2, column=HEADERS.index("Amount") + 1)
    assert amount_cell.number_format == '"$"#,##0.00'
    assert sheet.freeze_panes == "A2"


def test_html(funds, tmp_path):
    path = tmp_path / "funds.html"
    assert export_funds(funds, path) == 2

    page = path.read_text(encoding="utf-8")
    assert '<link rel="stylesheet" href="funds.css">' in page
    assert (tmp_path / "funds.css").exists()
    assert "2 funds, total $1,225.17" in page
    assert '<td class="number">$1,220.17</td>' in page
    assert "<td>2016-11-25</td>" in page
    assert "ŠIKIMIĆ ĐORĐE" in page
    # Jenkins blocks inline styles and scripts, so the page must not rely on them.
    assert "<style" not in page
    assert "<script" not in page


def test_html_escapes_values(funds, tmp_path):
    funds[0].owner_name = "<b>SMITH & SONS</b>"
    path = tmp_path / "funds.html"
    export_funds(funds, path)

    page = path.read_text(encoding="utf-8")
    assert "&lt;b&gt;SMITH &amp; SONS&lt;/b&gt;" in page


def test_unsupported_extension(funds, tmp_path):
    with pytest.raises(ValueError, match=r"use \.csv, \.xlsx, \.html"):
        export_funds(funds, tmp_path / "funds.pdf")
